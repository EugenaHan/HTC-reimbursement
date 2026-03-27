from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from pathlib import Path
from zipfile import ZipFile
import re


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = ROOT / "public" / "template.docx"

FULL_COL_WIDTHS_CM = [2.9] + [0.7] * 20


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("left", "top", "right", "bottom"):
        edge_data = kwargs.get(edge)
        if not edge_data:
            continue
        tag = "w:{}".format(edge)
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        for key, value in edge_data.items():
            element.set(qn("w:{}".format(key)), str(value))


def set_table_borders(table):
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(
                cell,
                top={"val": "single", "sz": "8", "color": "BFBFBF"},
                bottom={"val": "single", "sz": "8", "color": "BFBFBF"},
                left={"val": "single", "sz": "8", "color": "BFBFBF"},
                right={"val": "single", "sz": "8", "color": "BFBFBF"},
            )


def set_cell_text(cell, text, *, bold=False, size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "宋体"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def merge(row, start, end):
    cell = row.cells[start]
    for index in range(start + 1, end + 1):
        cell = cell.merge(row.cells[index])
    return cell


def apply_col_widths(table):
    table.autofit = False
    for row in table.rows:
        for index, width in enumerate(FULL_COL_WIDTHS_CM):
            row.cells[index].width = Cm(width)


def set_row_height(row, height_cm):
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    row.height = Cm(height_cm)


def postprocess_document_xml(docx_path):
    with ZipFile(docx_path, "r") as source:
        contents = {name: source.read(name) for name in source.namelist()}

    xml = contents["word/document.xml"].decode("utf-8")
    grid_xml = "".join(f'<w:gridCol w:w="{int(Cm(width).twips)}"/>' for width in FULL_COL_WIDTHS_CM)
    xml = re.sub(r"<w:tblGrid>.*?</w:tblGrid>", f"<w:tblGrid>{grid_xml}</w:tblGrid>", xml, count=1)
    xml = xml.replace('w:trHeight w:hRule="exact"', 'w:trHeight w:hRule="atLeast"')
    contents["word/document.xml"] = xml.encode("utf-8")

    with ZipFile(docx_path, "w") as target:
        for name, data in contents.items():
            target.writestr(name, data)


def build_document():
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.2)
    section.right_margin = Cm(1.2)

    title_1 = document.add_paragraph()
    title_1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_1.add_run("乾坤恒泰（北京）国际市场营销策划有限公司")
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(16)

    title_2 = document.add_paragraph()
    title_2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_2.add_run("{form_title}")
    run.bold = True
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(16)

    table = document.add_table(rows=17, cols=21)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    apply_col_widths(table)
    set_table_borders(table)

    row = table.rows[0]
    set_row_height(row, 1.0)
    set_cell_text(row.cells[0], "本表用途", size=12)
    set_cell_text(merge(row, 1, 10), "出差申请 {trip_checked}", size=10.5)
    set_cell_text(merge(row, 11, 20), "报销申请 {reimbursement_checked}", size=10.5)

    row = table.rows[1]
    set_row_height(row, 1.0)
    set_cell_text(row.cells[0], "员工姓名", size=12)
    set_cell_text(merge(row, 1, 8), "{name}", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(merge(row, 9, 13), "部门（HTC/SHOB）", size=12)
    set_cell_text(merge(row, 14, 20), "{dept}", align=WD_ALIGN_PARAGRAPH.LEFT)

    row = table.rows[2]
    set_row_height(row, 1.0)
    set_cell_text(row.cells[0], "出差时间", size=12)
    set_cell_text(merge(row, 1, 8), "（{start_date}）至（{end_date}）", size=12)
    set_cell_text(merge(row, 9, 13), "天数", size=12)
    set_cell_text(merge(row, 14, 20), "{trip_days}", align=WD_ALIGN_PARAGRAPH.LEFT)

    row = table.rows[3]
    set_row_height(row, 1.05)
    set_cell_text(row.cells[0], "出差事由", size=12)
    set_cell_text(merge(row, 1, 8), "{reason}", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(merge(row, 9, 13), "目的地", size=12)
    set_cell_text(merge(row, 14, 20), "{destination}", align=WD_ALIGN_PARAGRAPH.LEFT)

    row = table.rows[4]
    set_row_height(row, 1.1)
    set_cell_text(row.cells[0], "城市间交通方案", size=12)
    headers = ["出发时间", "抵达时间", "金额", "服务商", "报价时间"]
    for index, label in enumerate(headers):
        set_cell_text(merge(row, 1 + index * 4, 4 + index * 4), label, size=12)

    for offset, slot in enumerate((1, 2, 3), start=5):
        row = table.rows[offset]
        set_row_height(row, 1.15)
        set_cell_text(row.cells[0], f"方案{slot}：", size=12)
        set_cell_text(merge(row, 1, 4), f"{{transport_{slot}_departure}}", align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(merge(row, 5, 8), f"{{transport_{slot}_arrival}}", align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(merge(row, 9, 12), f"{{transport_{slot}_budget}}", align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(merge(row, 13, 16), f"{{transport_{slot}_vendor}}", align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(merge(row, 17, 20), f"{{transport_{slot}_quote_at}}", align=WD_ALIGN_PARAGRAPH.LEFT)

    row = table.rows[8]
    set_row_height(row, 1.1)
    set_cell_text(row.cells[0], "住宿方案", size=12)
    headers = ["入住时间", "退房时间", "金额", "服务商", "报价时间"]
    for index, label in enumerate(headers):
        set_cell_text(merge(row, 1 + index * 4, 4 + index * 4), label, size=12)

    for offset, slot in enumerate((1, 2, 3), start=9):
        row = table.rows[offset]
        set_row_height(row, 1.15)
        set_cell_text(row.cells[0], f"方案{slot}：", size=12)
        set_cell_text(merge(row, 1, 4), f"{{hotel_{slot}_check_in}}", align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(merge(row, 5, 8), f"{{hotel_{slot}_check_out}}", align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(merge(row, 9, 12), f"{{hotel_{slot}_budget}}", align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(merge(row, 13, 16), f"{{hotel_{slot}_vendor}}", align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell_text(merge(row, 17, 20), f"{{hotel_{slot}_quote_at}}", align=WD_ALIGN_PARAGRAPH.LEFT)

    row = table.rows[12]
    set_row_height(row, 1.05)
    set_cell_text(row.cells[0], "餐费总支出", size=12)
    set_cell_text(merge(row, 1, 8), "{meal_budget}", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(merge(row, 9, 15), "地面交通费总支出", size=12)
    set_cell_text(merge(row, 16, 20), "{ground_budget}", align=WD_ALIGN_PARAGRAPH.LEFT)

    row = table.rows[13]
    set_row_height(row, 1.1)
    set_cell_text(row.cells[0], "其它支出", size=12)
    set_cell_text(merge(row, 1, 20), "{other_budget}", align=WD_ALIGN_PARAGRAPH.LEFT)

    row = table.rows[14]
    set_row_height(row, 1.05)
    set_cell_text(row.cells[0], "总支出（预算）", size=12)
    set_cell_text(merge(row, 1, 8), "{total_budget}", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(merge(row, 9, 15), "总支出（实际）", size=12)
    set_cell_text(merge(row, 16, 20), "{total_actual}", align=WD_ALIGN_PARAGRAPH.LEFT)

    row = table.rows[15]
    set_row_height(row, 1.05)
    set_cell_text(row.cells[0], "部门主管审核", size=12)
    set_cell_text(merge(row, 1, 8), "", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(merge(row, 9, 15), "部门主管签名", size=12)
    set_cell_text(merge(row, 16, 20), "", align=WD_ALIGN_PARAGRAPH.LEFT)

    row = table.rows[16]
    set_row_height(row, 1.05)
    set_cell_text(row.cells[0], "总经理批准", size=12)
    set_cell_text(merge(row, 1, 8), "", align=WD_ALIGN_PARAGRAPH.LEFT)
    set_cell_text(merge(row, 9, 15), "总经理签名", size=12)
    set_cell_text(merge(row, 16, 20), "", align=WD_ALIGN_PARAGRAPH.LEFT)

    document.add_paragraph("")

    note_title = document.add_paragraph()
    note_title.paragraph_format.space_after = Pt(2)
    title_run = note_title.add_run("注意事项：")
    title_run.bold = True
    title_run.font.name = "宋体"
    title_run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    title_run.font.size = Pt(12)

    notes = [
        ("基本原则：", "厉行节约、预算控制、事前审批、真实合规", "。"),
        ("出差前填写申请表，明确", "出差事由、时间、地点、预算", "，并完成部门主管审核和总经理批准。"),
        ("城市间交通优先选择经济便捷的交通工具，", "三选一", "；特殊情况需提前说明，未经许可超标部分由个人承担。"),
        ("住宿标准：", "一线城市 1000 元/晚以内，省会城市 800 元/晚以内，其他城市 600 元/晚以内", "。"),
        ("餐费和市内交通费", "标准内根据有效票据实报实销", "；接待单位安排工作餐或交通工具的，不单独报销。"),
        ("报销须在", "出差结束后 7 个工作日内", "完成，超期需书面说明。"),
        ("所有票据及附件需", "真实、合法、完整", "，用于纳税和审计。"),
    ]

    for prefix, underline_text, suffix in notes:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.2
        run = paragraph.add_run(prefix)
        run.font.name = "宋体"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(12)

        if underline_text:
            underline_run = paragraph.add_run(underline_text)
            underline_run.underline = True
            underline_run.font.name = "宋体"
            underline_run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            underline_run.font.size = Pt(12)

        if suffix:
            suffix_run = paragraph.add_run(suffix)
            suffix_run.font.name = "宋体"
            suffix_run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            suffix_run.font.size = Pt(12)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)
    postprocess_document_xml(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
    print(OUTPUT_PATH)
